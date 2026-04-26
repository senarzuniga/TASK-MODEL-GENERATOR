
"""
Excel workbook builder.

Assembles a multi-sheet, professionally formatted Excel workbook containing
the B2B Commercial Task Model and an accompanying KPI Reference sheet.
"""

from __future__ import annotations

import itertools

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from .data import TASKS
from .formatter import (
    COLUMN_WIDTHS,
    THIN_BORDER,
    apply_cell_style,
    data_row_style,
    domain_group_style,
    header_style,
    maturity_fill,
    priority_fill,
)
from .models import Task


# ---------------------------------------------------------------------------
# Column indices (1-based) for columns that receive conditional formatting
# ---------------------------------------------------------------------------
_COL_PRIORITY = 11      # "Priority Level"
_COL_MATURITY = 17      # "Sales Maturity Level"


def _write_header_row(ws: Worksheet, headers: list[str]) -> None:
    """Write and style the header row."""
    ws.append(headers)
    style = header_style()
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        apply_cell_style(cell, style)


def _write_domain_separator(ws: Worksheet, domain: str, num_cols: int) -> None:
    """Insert a full-width domain group title row."""
    row_num = ws.max_row + 1
    ws.append([domain] + [""] * (num_cols - 1))
    style = domain_group_style()
    for col_idx in range(1, num_cols + 1):
        apply_cell_style(ws.cell(row=row_num, column=col_idx), style)
    # Merge across all columns for a clean group header
    ws.merge_cells(
        start_row=row_num,
        start_column=1,
        end_row=row_num,
        end_column=num_cols,
    )


def _write_task_row(ws: Worksheet, task: Task, row_index: int) -> None:
    """Write a single task row with full styling."""
    row_data = task.to_row()
    ws.append(row_data)
    current_row = ws.max_row
    base_style = data_row_style(row_index)

    for col_idx, value in enumerate(row_data, start=1):
        cell = ws.cell(row=current_row, column=col_idx)
        apply_cell_style(cell, base_style)

        # Override fill for priority and maturity columns
        if col_idx == _COL_PRIORITY:
            # Apply specific fill and font styles for priority column
            cell.fill = priority_fill(str(value))
            cell.font = Font(bold=True, size=9, name="Calibri")
            cell.alignment = Alignment(horizontal="center", vertical="top")
        elif col_idx == _COL_MATURITY:
            # Apply specific fill and font styles for maturity column
            cell.fill = maturity_fill(str(value))
            cell.font = Font(bold=True, size=9, name="Calibri")
            cell.alignment = Alignment(horizontal="center", vertical="top")


def _set_column_widths(ws: Worksheet) -> None:
    """Apply the defined column widths to the worksheet."""
    for col_idx, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _freeze_header(ws: Worksheet) -> None:
    """Freeze the first row and first two columns."""
    ws.freeze_panes = "C2"


def _add_auto_filter(ws: Worksheet, headers: list[str]) -> None:
    """Add Excel AutoFilter to the header row."""
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col}1"


def _build_task_sheet(wb: Workbook) -> Worksheet:
    """Build and return the main B2B Task Model worksheet."""
    ws = wb.active
    ws.title = "B2B Task Model"

    headers = Task.headers()
    _write_header_row(ws, headers)

    # Group tasks by domain and write domain separators
    data_row_counter = 0
    for domain, group in itertools.groupby(TASKS, key=lambda t: t.domain):
        _write_domain_separator(ws, domain, len(headers))
        for task in group:
            _write_task_row(ws, task, data_row_counter)
            data_row_counter += 1

    _set_column_widths(ws)
    _freeze_header(ws)
    _add_auto_filter(ws, headers)

    # Row height for header
    ws.row_dimensions[1].height = 40

    return ws


def _build_kpi_reference_sheet(wb: Workbook) -> Worksheet:
    """Build a KPI reference sheet listing all KPIs and their descriptions."""
    ws = wb.create_sheet(title="KPI Reference")

    kpi_headers = ["KPI Name", "Definition", "How to Measure", "Target Benchmark"]
    kpis = [
        [
            "Forecast Accuracy",
            "How closely predicted revenue matches actual revenue",
            "(1 − |Actual − Forecast| / Forecast) × 100%",
            "> 90%",
        ],
        [
            "Win Rate",
            "% of qualified opportunities converted to closed-won",
            "Closed-Won / (Closed-Won + Closed-Lost) × 100%",
            "> 30% (industry avg)",
        ],
        [
            "Sales Cycle Length",
            "Average calendar days from opportunity creation to close",
            "Mean of (Close Date − Create Date) for closed-won deals",
            "Trending downward YoY",
        ],
        [
            "Pipeline Coverage Ratio",
            "Total pipeline value as multiple of quota",
            "Open Pipeline Value / Quota",
            "> 3×",
        ],
        [
            "Account Revenue Growth",
            "YoY revenue increase per key account",
            "(Current Year Revenue − Prior Year Revenue) / Prior Year Revenue × 100%",
            "> 10% per annum",
        ],
        [
            "Customer Retention",
            "% of customers renewing / not churning within period",
            "Retained Customers / Total Customers at Period Start × 100%",
            "> 90%",
        ],
        [
            "Upsell Revenue",
            "Incremental revenue from selling additional products to existing customers",
            "Sum of upsell opportunity closed-won value",
            "> 20% of total revenue",
        ],
        [
            "Sales Productivity",
            "Revenue generated per sales FTE",
            "Total Revenue / Number of Sales FTEs",
            "Trending upward YoY",
        ],
        [
            "Churn Rate",
            "% of customers lost within a period",
            "Lost Customers / Customers at Period Start × 100%",
            "< 5% annually",
        ],
        [
            "Net Revenue Retention",
            "Revenue retained + expansion, minus churn and downsell",
            "(Starting MRR + Expansion − Downsell − Churn) / Starting MRR × 100%",
            "> 110%",
        ],
        [
            "Lead-to-Opportunity Conversion",
            "% of leads that become qualified opportunities",
            "Qualified Opportunities Created / Total Leads × 100%",
            "> 20%",
        ],
        [
            "Meetings Booked",
            "Number of qualified discovery/demo meetings per SDR per week",
            "Count of meetings booked in CRM",
            "> 5 per SDR / week",
        ],
        [
            "Revenue Attainment",
            "% of revenue target achieved in the period",
            "Actual Revenue / Revenue Target × 100%",
            "> 100%",
        ],
    ]

    # Header row
    ws.append(kpi_headers)
    h_style = header_style()
    for col_idx in range(1, len(kpi_headers) + 1):
        apply_cell_style(ws.cell(row=1, column=col_idx), h_style)

    # Data rows
    kpi_col_widths = [30, 55, 55, 25]
    for row_idx, kpi in enumerate(kpis, start=1):
        ws.append(kpi)
        row_style = data_row_style(row_idx)
        for col_idx in range(1, len(kpi_headers) + 1):
            apply_cell_style(ws.cell(row=row_idx + 1, column=col_idx), row_style)

    for col_idx, width in enumerate(kpi_col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    return ws


def _build_legend_sheet(wb: Workbook) -> Worksheet:
    """Build a legend / guide sheet explaining column values and colour coding."""
    ws = wb.create_sheet
