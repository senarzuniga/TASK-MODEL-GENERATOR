
"""
Excel formatting utilities.

Provides colour palettes, cell styles, and helper functions used by the
workbook builder to produce a professional, print-ready output.
"""

from __future__ import annotations

from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    GradientFill,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Colour palette (hex without leading #)
# ---------------------------------------------------------------------------
COLOURS = {
    # Header row
    "header_bg": "1F3864",       # Dark navy
    "header_fg": "FFFFFF",       # White text

    # Domain group title rows
    "domain_bg": "2E75B6",       # Mid blue
    "domain_fg": "FFFFFF",

    # Data row alternates
    "row_even": "EBF3FB",        # Very light blue
    "row_odd": "FFFFFF",         # White

    # Priority highlights
    "priority_high": "FF4C4C",   # Red
    "priority_medium": "FFA500", # Orange
    "priority_low": "70AD47",    # Green

    # Maturity level
    "maturity_basic": "D9D9D9",
    "maturity_advanced": "FFE699",
    "maturity_elite": "C6EFCE",

    # Border
    "border": "BDD7EE",
}

# ---------------------------------------------------------------------------
# Thin border style
# ---------------------------------------------------------------------------
_thin_side = Side(style="thin", color=COLOURS["border"])
THIN_BORDER = Border(
    left=_thin_side,
    right=_thin_side,
    top=_thin_side,
    bottom=_thin_side,
)

_medium_side = Side(style="medium", color=COLOURS["header_bg"])
HEADER_BORDER = Border(
    left=_medium_side,
    right=_medium_side,
    top=_medium_side,
    bottom=_medium_side,
)


# ---------------------------------------------------------------------------
# Style factories
# ---------------------------------------------------------------------------

def header_style() -> dict:
    """Return keyword arguments for a header cell."""
    return {
        "font": Font(
            bold=True,
            color=COLOURS["header_fg"],
            size=10,
            name="Calibri",
        ),
        "fill": PatternFill(
            fill_type="solid",
            fgColor=COLOURS["header_bg"],
        ),
        "alignment": Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        ),
        "border": HEADER_BORDER,
    }


def domain_group_style() -> dict:
    """Return keyword arguments for a domain group separator row."""
    return {
        "font": Font(
            bold=True,
            color=COLOURS["domain_fg"],
            size=10,
            name="Calibri",
            italic=True,
        ),
        "fill": PatternFill(
            fill_type="solid",
            fgColor=COLOURS["domain_bg"],
        ),
        "alignment": Alignment(
            horizontal="left",
            vertical="center",
        ),
        "border": THIN_BORDER,
    }


def data_row_style(row_index: int) -> dict:
    """Return keyword arguments for a standard data cell."""
    # Alternate row colors for better readability
    bg = COLOURS["row_even"] if row_index % 2 == 0 else COLOURS["row_odd"]
    return {
        "font": Font(size=9, name="Calibri"),
        "fill": PatternFill(fill_type="solid", fgColor=bg),
        "alignment": Alignment(
            vertical="top",
            wrap_text=True,
        ),
        "border": THIN_BORDER,
    }


def priority_fill(priority_value: str) -> PatternFill:
    """Return a PatternFill based on priority text."""
    # Map priority levels to specific colors
    mapping = {
        "High": COLOURS["priority_high"],
        "Medium": COLOURS["priority_medium"],
        "Low": COLOURS["priority_low"],
    }
    colour = mapping.get(priority_value, COLOURS["row_odd"])
    return PatternFill(fill_type="solid", fgColor=colour)


def maturity_fill(maturity_value: str) -> PatternFill:
    """Return a PatternFill based on maturity level text."""
    # Map maturity levels to specific colors
    mapping = {
        "Basic": COLOURS["maturity_basic"],
        "Advanced": COLOURS["maturity_advanced"],
        "Elite": COLOURS["maturity_elite"],
    }
    colour = mapping.get(maturity_value, COLOURS["row_odd"])
    return PatternFill(fill_type="solid", fgColor=colour)


# ---------------------------------------------------------------------------
# Column width configuration
# ---------------------------------------------------------------------------
# (column index 1-based, width in characters)
COLUMN_WIDTHS: dict[int, float] = {
    1:  18,   # Domain
    2:  28,   # Task Name
    3:  38,   # Objective
    4:  30,   # Trigger
    5:  50,   # Execution Method
    6:  35,   # Key Inputs
    7:  35,   # Expected Output
    8:  28,   # Primary KPI
    9:  35,   # Secondary KPIs
    10: 16,   # Frequency
    11: 14,   # Priority
    12: 22,   # Owner
    13: 28,   # Dependencies
    14: 30,   # Tools
    15: 40,   # AI Role
    16: 20,   # Automation Level
    17: 22,   # Maturity Level
}


def apply_cell_style(cell, style: dict) -> None:
    """Apply a style dict to an openpyxl Cell object."""
    if "font" in style:
        cell.font = style["font"]
    if "fill" in style:
        cell.fill = style["fill"]
    if "alignment" in style:
        cell.alignment = style["alignment"]
    if "border" in style:
        cell.border = style["border"]

