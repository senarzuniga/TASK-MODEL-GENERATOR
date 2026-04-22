
"""
Tests for the B2B Commercial Task Model Generator.
"""

import os
import tempfile
import logging

import pytest
from openpyxl import load_workbook

from generator import build_workbook
from generator.data import TASKS
from generator.models import AutomationLevel, MaturityLevel, Priority, Task

# Configure logger for testing
logging.basicConfig(level=logging.DEBUG, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Model tests
# ---------------------------------------------------------------------------

class TestTask:
    def test_headers_length(self):
        assert len(Task.headers()) == 17

    def test_to_row_length_matches_headers(self):
        task = TASKS[0]
        assert len(task.to_row()) == len(Task.headers())

    def test_to_row_values_are_strings(self):
        for task in TASKS:
            for value in task.to_row():
                assert isinstance(value, str), (
                    f"Expected str, got {type(value)} in task '{task.name}'"
                )

    def test_enum_values_present_in_row(self):
        task = TASKS[0]
        row = task.to_row()
        assert task.priority.value in row
        assert task.automation_level.value in row
        assert task.maturity_level.value in row


class TestEnums:
    def test_automation_levels(self):
        assert AutomationLevel.MANUAL.value == "Manual"
        assert AutomationLevel.SEMI.value == "Semi"
        assert AutomationLevel.FULL.value == "Full"

    def test_maturity_levels(self):
        assert MaturityLevel.BASIC.value == "Basic"
        assert MaturityLevel.ADVANCED.value == "Advanced"
        assert MaturityLevel.ELITE.value == "Elite"

    def test_priority_levels(self):
        assert Priority.HIGH.value == "High"
        assert Priority.MEDIUM.value == "Medium"
        assert Priority.LOW.value == "Low"


# ---------------------------------------------------------------------------
# Data catalogue tests
# ---------------------------------------------------------------------------

class TestTaskCatalogue:
    def test_minimum_task_count(self):
        """Catalogue must contain at least 20 tasks."""
        assert len(TASKS) >= 20

    def test_all_required_domains_present(self):
        domains = {t.domain for t in TASKS}
        required = {
            "Pipeline Management",
            "KAM",
            "Customer Growth",
            "Sales Execution",
            "After Sales",
            "Strategy & Performance",
            "Prospecting",
        }
        assert required.issubset(domains), (
            f"Missing domains: {required - domains}"
        )

    def test_no_empty_required_fields(self):
        required_fields = [
            "domain", "name", "objective", "trigger", "method",
            "inputs", "output", "primary_kpi", "owner",
        ]
        for task in TASKS:
            for field in required_fields:
                value = getattr(task, field)
                assert value and value.strip(), (
                    f"Empty field '{field}' in task '{task.name}'"
                )

    def test_valid_enum_assignments(self):
        for task in TASKS:
            assert isinstance(task.priority, Priority)
            assert isinstance(task.automation_level, AutomationLevel)
            assert isinstance(task.maturity_level, MaturityLevel)

    def test_unique_task_names(self):
        names = [t.name for t in TASKS]
        assert len(names) == len(set(names)), "Duplicate task names found"


# ---------------------------------------------------------------------------
# Workbook builder tests
# ---------------------------------------------------------------------------

class TestBuildWorkbook:
    def test_creates_file(self, tmp_path):
        out = str(tmp_path / "test_output.xlsx")
        result = build_workbook(out)
        assert result == out
        assert os.path.isfile(out)

    def test_three_sheets(self, tmp_path):
        out = str(tmp_path / "wb.xlsx")
        build_workbook(out)
        wb = load_workbook(out)
        assert wb.sheetnames == ["B2B Task Model", "KPI Reference", "Legend & Guide"]

    def test_task_sheet_header(self, tmp_path):
        out = str(tmp_path / "wb.xlsx")
        build_workbook(out)
        wb = load_workbook(out)
        ws = wb["B2B Task Model"]
        headers = [ws.cell(1, c).value for c in range(1, 18)]
        assert headers == Task.headers()

    def test_task_sheet_has_correct_column_count(self, tmp_path):
        out = str(tmp_path / "wb.xlsx")
        build_workbook(out)
        wb = load_workbook(out)
        ws = wb["B2B Task Model"]
        assert ws.max_column == 17

    def test_task_sheet_data_rows_match_catalogue(self, tmp_path):
        out = str(tmp_path / "wb.xlsx")
        build_workbook(out)
        wb = load_workbook(out)
        ws = wb["B2B Task Model"]
        # Count rows where column 2 (Task Name) is populated
        data_rows = sum(
            1 for row in ws.iter_rows(min_row=2)
            if row[1].value  # column B = Task Name
        )
        assert data_rows == len(TASKS)

    def test_kpi_reference_sheet_has_data(self, tmp_path):
        out = str(tmp_path / "wb.xlsx")
        build_workbook(out)
        wb = load_workbook(out)
        ws = wb["KPI Reference"]
        # Header + at least 10 KPI rows
        assert ws.max_row >= 11

    def test_legend_sheet_has_data(self, tmp_path):
        out = str(tmp_path / "wb.xlsx")
        build_workbook(out)
        wb = load_workbook(out)
        ws = wb["Legend & Guide"]
        assert ws.max_row >= 20

    def test_freeze_panes_set(self, tmp_path):
        out = str(tmp_path / "wb.xlsx")
        build_workbook(out)
        wb = load_workbook(out)
        ws = wb["B2B Task Model"]
        assert ws.freeze_panes == "C2"

    def test_autofilter_set(self, tmp_path):
        out = str(tmp_path / "wb.xlsx")
        build_workbook(out)
        wb = load_workbook(out)
        ws = wb["B2B Task Model"]
        assert ws.auto_filter.ref is not None

    def test_logging_on_error(self, caplog):
        with caplog.at_level(logging.ERROR):
            try:
                raise ValueError("Test error")
            except ValueError as e:
                logger.error("Caught an error: %s", e)
        assert "Caught an error: Test error" in caplog.text


# ---------------------------------------------------------------------------
# CLI entry point tests
# ---------------------------------------------------------------------------

class TestCLI:
    def test_main_returns_zero_on_success(self, tmp_path):
        from main import main
        out = str(tmp_path / "cli_test.xlsx")
        assert main(["--output", out]) == 0

    def test_main_creates_file(self, tmp_path):
        from main import main
        out = str(tmp_path / "cli_test.xlsx")
        main(["--output", out])
        assert os.path.isfile(out)

    def test_main_default_output(self, tmp_path, monkeypatch):
        from main import main
        monkeypatch.chdir(tmp_path)
        result = main([])
        assert result == 0
        assert os.path.isfile(tmp_path / "Elite_B2B_Sales_Task_Model.xlsx")

