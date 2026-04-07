"""Excel workbook generator for Task Models."""

from __future__ import annotations

import io
from pathlib import Path
from typing import List, Union

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import TaskModel


# ── Colour palette ───────────────────────────────────────────────────────────
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
SUBHEADER_FILL = PatternFill("solid", fgColor="2E75B6")
ALT_ROW_FILL = PatternFill("solid", fgColor="D6E4F0")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
SUBHEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Calibri", size=10)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _apply_header(cell, text: str, fill=HEADER_FILL, font=HEADER_FONT) -> None:
    cell.value = text
    cell.fill = fill
    cell.font = font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = THIN_BORDER


def _auto_fit_columns(ws, min_width: int = 12, max_width: int = 50) -> None:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_width), max_width)


class WorkbookGenerator:
    """Generates an Excel workbook from a list of TaskModels."""

    def __init__(self, task_models: List[TaskModel]) -> None:
        self.task_models = task_models

    def generate(self, output_path: Union[str, Path, io.IOBase]) -> Union[Path, io.IOBase]:
        """Build the workbook and write it to *output_path* (file path or buffer)."""
        is_buffer = isinstance(output_path, (io.RawIOBase, io.BufferedIOBase))
        if not is_buffer:
            output_path = Path(output_path)
            output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default sheet

        self._build_summary_sheet(wb)
        for model in self.task_models:
            self._build_model_sheet(wb, model)

        wb.save(output_path)
        return output_path

    # ── Summary sheet ─────────────────────────────────────────────────────────
    def _build_summary_sheet(self, wb: openpyxl.Workbook) -> None:
        ws = wb.create_sheet("Summary")
        ws.freeze_panes = "A2"

        headers = ["ID", "Name", "Category", "Priority", "Status", "Owner", "Steps", "Est. Hours"]
        for col_idx, header in enumerate(headers, start=1):
            _apply_header(ws.cell(row=1, column=col_idx), header)

        for row_idx, model in enumerate(self.task_models, start=2):
            fill = ALT_ROW_FILL if row_idx % 2 == 0 else PatternFill()
            values = [
                model.id,
                model.name,
                model.category,
                model.priority,
                model.status,
                model.owner,
                model.step_count,
                model.total_estimated_hours,
            ]
            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = BODY_FONT
                cell.fill = fill
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="center")

        _auto_fit_columns(ws)

    # ── Per-model sheet ───────────────────────────────────────────────────────
    def _build_model_sheet(self, wb: openpyxl.Workbook, model: TaskModel) -> None:
        sheet_name = model.name[:31]  # Excel sheet name limit
        ws = wb.create_sheet(sheet_name)
        ws.freeze_panes = "A4"

        # Title row
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = f"Task Model: {model.name}"
        title_cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=13)
        title_cell.fill = HEADER_FILL
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        # Metadata row
        meta = [
            ("ID", model.id),
            ("Category", model.category),
            ("Priority", model.priority),
            ("Status", model.status),
            ("Owner", model.owner),
            ("Est. Hours", model.total_estimated_hours),
        ]
        ws.merge_cells("A2:G2")
        meta_cell = ws["A2"]
        meta_cell.value = "  |  ".join(f"{k}: {v}" for k, v in meta)
        meta_cell.font = Font(name="Calibri", italic=True, size=9, color="FFFFFF")
        meta_cell.fill = SUBHEADER_FILL
        meta_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 18

        # Steps header
        step_headers = ["#", "Step Name", "Description", "Assignee", "Est. Hours", "Status"]
        for col_idx, header in enumerate(step_headers, start=1):
            _apply_header(ws.cell(row=3, column=col_idx), header, fill=SUBHEADER_FILL, font=SUBHEADER_FONT)
        ws.row_dimensions[3].height = 20

        # Steps data
        for row_idx, step in enumerate(model.steps, start=4):
            fill = ALT_ROW_FILL if row_idx % 2 == 0 else PatternFill()
            values = [
                step.order,
                step.name,
                step.description,
                step.assignee,
                step.estimated_hours,
                step.status,
            ]
            for col_idx, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = BODY_FONT
                cell.fill = fill
                cell.border = THIN_BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=(col_idx == 3))

        # Description section
        if model.description:
            desc_row = len(model.steps) + 5
            ws.cell(row=desc_row, column=1, value="Description:").font = Font(bold=True, size=10)
            ws.merge_cells(f"B{desc_row}:G{desc_row}")
            ws.cell(row=desc_row, column=2, value=model.description).alignment = Alignment(wrap_text=True)

        _auto_fit_columns(ws)
